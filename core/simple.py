"""The six-section management interface. Legacy workflows remain stored unchanged."""
import csv
import io
from datetime import date
from decimal import Decimal
from django import forms
from django.contrib import messages
from django.db import IntegrityError
from django.db.models import Q
from django.http import FileResponse, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST
from .models import (FinanceTransaction, TransactionRevision, SavedScenario, SimpleDecision,
                     DecisionVote, DecisionComment, Document, DocumentLink, Journal, Decision, CashOpening, CashAccount, ReinvestmentPlan)
from .finance import BusinessError, money
from .simple_finance import calculate_scenario, outstanding, summary, account_balances
from .security import require_role, require_reauth
from .views import workspace_view, context
from . import services, reports

NAV=[('dashboard','Dashboard','◈'),('transactions','Transactions','↔'),('calculator','Decision Calculator','◇'),
     ('decisions','Decisions & Approvals','✓'),('documents','Documents','▱'),('reports','Reports','▥')]

def ctx(request,div,section,**kwargs):
    data=context(request,div,section=section,title=dict((x[0],x[1]) for x in NAV)[section],**kwargs)
    data['nav']=NAV
    return data

def files(request,div,parent=None):
    for item in request.FILES.getlist('files'):
        doc=services.upload(div,request.user,item.name,item.content_type,item.read(10*1024*1024+1))
        if parent:
            source=next(iter(parent.values()))
            DocumentLink.objects.create(document=doc,category=getattr(source,'category','')[:80],
                supplier=getattr(source,'party','')[:200],reference=getattr(source,'reference','')[:100],**parent)

class TransactionForm(forms.ModelForm):
    class Meta:
        model=FinanceTransaction
        fields=['date','direction','amount','total_due','category','description','party','method','account','reinvestment_plan','reference','notes','treatment','settles','refund_of','decision']
        widgets={'date':forms.DateInput(attrs={'type':'date'}),'notes':forms.Textarea(attrs={'rows':3})}
        labels={'direction':'Payment type','amount':'Amount paid or received (SAR)','party':'Payer or payee (optional)','method':'Payment method','reference':'Reference number (optional)','total_due':'Total owed (optional)','settles':'Partial payment for (optional)','refund_of':'Refund of (optional)','treatment':'Effect on estimated profit'}
    def __init__(self,*args,division=None,**kwargs):
        super().__init__(*args,**kwargs)
        self.division=division
        self.fields['settles'].queryset=FinanceTransaction.objects.filter(division=division,total_due__isnull=False,settles__isnull=True,refund_of__isnull=True)
        self.fields['refund_of'].queryset=FinanceTransaction.objects.filter(division=division,refund_of__isnull=True)
        self.fields['decision'].queryset=SimpleDecision.objects.filter(division=division)
        self.fields['account'].queryset=CashAccount.objects.filter(division=division)
        self.fields['reinvestment_plan'].queryset=ReinvestmentPlan.objects.filter(division=division)
        if self.instance._state.adding: self.fields['date'].initial=timezone.localdate()
    def clean(self):
        p=super().clean(); amount=p.get('amount'); due=p.get('total_due'); settles=p.get('settles'); refund=p.get('refund_of')
        if amount is None: return p
        if p.get('category','').strip().lower() in ('capital contribution','capital repayment','partner distribution','mold payment','production payment') and p.get('treatment') in ('income','expense'):
            self.add_error('treatment','This cash category does not create sales income or operating expense.')
        if p.get('account') and p.get('method')!=p['account'].kind: self.add_error('account','Choose an account matching the payment method.')
        if p.get('reinvestment_plan') and p.get('direction')!='paid': self.add_error('reinvestment_plan','Reinvestment allocations can only be linked to spending.')
        if due is not None and amount>due: self.add_error('amount','Money received or paid cannot exceed the total owed.')
        if settles and refund: raise forms.ValidationError('Choose either a partial payment or a refund.')
        if settles:
            if p.get('direction')!=settles.direction: self.add_error('direction','A partial payment must have the same money direction.')
            if due is not None: self.add_error('total_due','The original entry holds the total owed.')
            if amount<=0: self.add_error('amount','Enter a payment greater than zero.')
            allowance=outstanding(settles)+(self.instance.amount if not self.instance._state.adding and self.instance.settles_id==settles.pk else 0)
            if amount>allowance: self.add_error('amount','Payment exceeds the remaining balance.')
            p['treatment']='none'
        if refund:
            if p.get('direction')==refund.direction: self.add_error('direction','A refund reverses the original money direction.')
            if due is not None: self.add_error('total_due','A refund cannot create a new amount owed.')
            if amount<=0: self.add_error('amount','Enter a refund greater than zero.')
            already=sum((x.amount for x in refund.refunds.exclude(pk=self.instance.pk if not self.instance._state.adding else None)),Decimal(0))
            if amount+already>refund.amount: self.add_error('amount','Refunds cannot exceed the original cash amount.')
            p['treatment']=refund.treatment
        if not self.instance._state.adding:
            if settles and settles.pk==self.instance.pk: self.add_error('settles','An entry cannot settle itself.')
            if refund and refund.pk==self.instance.pk: self.add_error('refund_of','An entry cannot refund itself.')
            if self.instance.payments.exists() or self.instance.refunds.exists():
                protected=['amount','total_due','direction','treatment','settles','refund_of']
                if any(getattr(self.instance,k)!=p.get(k) for k in protected):
                    raise forms.ValidationError('This entry has linked payments or refunds. Add a correction instead of changing its amounts.')
        if not p.get('reference') and p.get('date') and p.get('category') and p.get('description'):
            duplicate=FinanceTransaction.objects.filter(division=self.division,date=p['date'],direction=p.get('direction'),amount=amount,category__iexact=p['category'],description__iexact=p['description'],party__iexact=p.get('party',''))
            if not self.instance._state.adding: duplicate=duplicate.exclude(pk=self.instance.pk)
            if duplicate.exists(): raise forms.ValidationError('This looks like a duplicate. Open the existing entry or add a distinct reference.')
        return p

def snapshot(t):
    return {k:str(getattr(t,k)) if getattr(t,k) is not None else None for k in ('date','direction','amount','total_due','category','description','party','method','account_id','reinvestment_plan_id','reference','notes','treatment','recognized_amount','settles_id','refund_of_id','decision_id')}

def save_transaction(form,div,user):
    services.lock(div); t=form.save(commit=False)
    previous=snapshot(FinanceTransaction.objects.get(pk=t.pk,division=div)) if not t._state.adding else None
    t.division=div
    if t._state.adding: t.created_by=user
    if t.settles_id: t.recognized_amount=0
    elif t.refund_of_id: t.recognized_amount=-t.amount if t.treatment in ('income','expense') else 0
    elif t.treatment in ('income','expense'): t.recognized_amount=t.total_due if t.total_due is not None else t.amount
    else: t.recognized_amount=0
    try: t.save()
    except IntegrityError: raise BusinessError('This reference already exists in this workspace.')
    TransactionRevision.objects.create(transaction=t,actor=user,snapshot=snapshot(t))
    services.audit(div,user,'simple_transaction.updated' if previous else 'simple_transaction.created',t.pk,
                   {'before':previous,'after':snapshot(t)})
    return t

def filters(request):
    start=request.GET.get('from','').strip(); end=request.GET.get('to','').strip(); category=request.GET.get('category','').strip()
    try: return (date.fromisoformat(start) if start else None,date.fromisoformat(end) if end else None,category)
    except ValueError: raise BusinessError('Use valid dates for the report filter.')

def filtered_entries(div,request):
    start,end,category=filters(request)
    qs=FinanceTransaction.objects.filter(division=div).select_related('created_by').order_by('-date','-created_at')
    if start: qs=qs.filter(date__gte=start)
    if end: qs=qs.filter(date__lte=end)
    if category: qs=qs.filter(category=category)
    q=request.GET.get('q','').strip()[:100]
    if q: qs=qs.filter(Q(description__icontains=q)|Q(party__icontains=q)|Q(reference__icontains=q)|Q(category__icontains=q))
    return qs

@workspace_view
def dashboard(request,div):
    start,end,category=filters(request); data=summary(div,start,end,category)
    return render(request,'core/simple_dashboard.html',ctx(request,div,'dashboard',summary=data,balances=account_balances(div),
        recent=filtered_entries(div,request)[:8],pending=SimpleDecision.objects.filter(division=div,status='pending')[:6],
        categories=FinanceTransaction.objects.filter(division=div).order_by('category').values_list('category',flat=True).distinct()))

@require_POST
@workspace_view
def opening_balance(request,div):
    require_role(request.user,div,('operator',)); services.lock(div)
    method=request.POST.get('method')
    if method not in ('bank','cash'): raise BusinessError('Choose bank or cash.')
    try: as_of=date.fromisoformat(request.POST.get('as_of',''))
    except ValueError: raise BusinessError('Enter a valid balance date.')
    amount=money(request.POST.get('amount',''))
    if amount<0: raise BusinessError('An opening balance cannot be negative.')
    opening,created=CashOpening.objects.get_or_create(division=div,method=method,defaults={'amount':amount,'as_of':as_of,'recorded_by':request.user})
    if not created: raise BusinessError('Opening balance already exists. Keep the original and record a correction transaction.')
    services.audit(div,request.user,'cash_opening.recorded',opening.pk,{'method':method,'amount':str(amount),'as_of':str(as_of)})
    return redirect('dashboard',division_id=div.pk)

@workspace_view
def transactions(request,div):
    return render(request,'core/simple_transactions.html',ctx(request,div,'transactions',entries=filtered_entries(div,request)[:200],
        categories=FinanceTransaction.objects.filter(division=div).order_by('category').values_list('category',flat=True).distinct(),summary=summary(div,*filters(request))))

@workspace_view
def transaction_form(request,div,entry_id=None):
    require_role(request.user,div,('operator',))
    t=get_object_or_404(FinanceTransaction,division=div,pk=entry_id) if entry_id else None
    if t and (t.mold_id or t.batch_id or t.invoice_id or hasattr(t,'capital_contribution') or hasattr(t,'capital_repayment')):
        raise BusinessError('Linked payments are fixed. Record a refund or correction against the original record.')
    form=TransactionForm(request.POST or None,instance=t,division=div)
    if request.method=='POST' and form.is_valid():
        t=save_transaction(form,div,request.user); files(request,div,{'transaction':t})
        messages.success(request,'Transaction saved with an audit entry.')
        return redirect('simple_transaction_detail',division_id=div.pk,entry_id=t.pk)
    return render(request,'core/simple_form.html',ctx(request,div,'transactions',form=form,heading='Edit transaction' if t else 'Quick Add Payment',upload=True,entry=t))

@workspace_view
def transaction_detail(request,div,entry_id):
    t=get_object_or_404(FinanceTransaction,division=div,pk=entry_id)
    linked=bool(t.mold_id or t.batch_id or t.invoice_id or hasattr(t,'capital_contribution') or hasattr(t,'capital_repayment'))
    if request.method=='POST':
        require_role(request.user,div,('operator',))
        if not request.FILES.getlist('files'): raise BusinessError('Choose at least one file.')
        files(request,div,{'transaction':t})
        return redirect('simple_transaction_detail',division_id=div.pk,entry_id=t.pk)
    return render(request,'core/simple_transaction_detail.html',ctx(request,div,'transactions',entry=t,balance=outstanding(t),
        linked=linked,attachments=t.attachments.select_related('document'),revisions=t.revisions.select_related('actor').order_by('-created_at')))

class ScenarioForm(forms.ModelForm):
    class Meta:
        model=SavedScenario
        fields=['name','investment','unit_cost','selling_price','quantity','extra_expenses']
        labels={'investment':'Investment / one-time cost (SAR)','unit_cost':'Cost per unit (SAR)','selling_price':'Selling price per unit (SAR)','extra_expenses':'Other fixed expenses (SAR)'}
    def clean(self):
        p=super().clean()
        if all(p.get(k) is not None for k in ('investment','unit_cost','selling_price','quantity','extra_expenses')):
            calculate_scenario(p['investment'],p['unit_cost'],p['selling_price'],p['quantity'],p['extra_expenses'])
        return p

@workspace_view
def calculator(request,div):
    form=ScenarioForm(request.POST or None)
    if request.method=='POST':
        require_role(request.user,div,('operator',))
        if form.is_valid():
            services.lock(div); s=form.save(commit=False); s.division=div; s.created_by=request.user; s.save()
            services.audit(div,request.user,'scenario.saved',s.pk,{'name':s.name})
            return redirect('simple_calculator',division_id=div.pk)
    scenarios=[]
    for s in SavedScenario.objects.filter(division=div).order_by('-created_at')[:30]:
        scenarios.append((s,calculate_scenario(s.investment,s.unit_cost,s.selling_price,s.quantity,s.extra_expenses)))
    return render(request,'core/simple_calculator.html',ctx(request,div,'calculator',form=form,scenarios=scenarios))

class DecisionForm(forms.ModelForm):
    class Meta:
        model=SimpleDecision
        fields=['title','description','estimated_cost','expected_benefit','risks','scenario','required_approvals']
        widgets={'description':forms.Textarea(attrs={'rows':4}),'expected_benefit':forms.Textarea(attrs={'rows':2}),'risks':forms.Textarea(attrs={'rows':2})}
        labels={'required_approvals':'Independent approvals needed'}
    def __init__(self,*args,division=None,**kwargs):
        super().__init__(*args,**kwargs)
        self.fields['scenario'].queryset=SavedScenario.objects.filter(division=division)

@workspace_view
def decisions(request,div):
    return render(request,'core/simple_decisions.html',ctx(request,div,'decisions',
        decisions=SimpleDecision.objects.filter(division=div).select_related('created_by').order_by('-created_at'),
        legacy=Decision.objects.filter(division=div).order_by('-created_at')[:30]))

@workspace_view
def decision_new(request,div):
    require_role(request.user,div,('operator',))
    form=DecisionForm(request.POST or None,division=div)
    if request.method=='POST' and form.is_valid():
        services.lock(div); d=form.save(commit=False); d.division=div; d.created_by=request.user; d.save()
        services.audit(div,request.user,'simple_decision.created',d.pk,{'required_approvals':d.required_approvals})
        files(request,div,{'decision':d})
        return redirect('simple_decision_detail',division_id=div.pk,decision_id=d.pk)
    return render(request,'core/simple_form.html',ctx(request,div,'decisions',form=form,heading='Quick Add Decision',upload=True))

@workspace_view
def decision_detail(request,div,decision_id):
    d=get_object_or_404(SimpleDecision,division=div,pk=decision_id)
    return render(request,'core/simple_decision_detail.html',ctx(request,div,'decisions',decision=d,
        votes=d.votes.select_related('user'),comments=d.comments.select_related('user').order_by('created_at'),
        attachments=d.attachments.select_related('document'),linked=d.transactions.all(),
        result=calculate_scenario(d.scenario.investment,d.scenario.unit_cost,d.scenario.selling_price,d.scenario.quantity,d.scenario.extra_expenses) if d.scenario_id else None))

@require_POST
@workspace_view
def decision_action(request,div,decision_id):
    require_role(request.user,div,('operator',)); services.lock(div)
    d=SimpleDecision.objects.select_for_update().get(division=div,pk=decision_id)
    action=request.POST.get('action')
    if action=='comment':
        text=request.POST.get('text','').strip()
        if not text or len(text)>3000: raise BusinessError('Enter a comment of up to 3,000 characters.')
        DecisionComment.objects.create(decision=d,user=request.user,text=text)
        services.audit(div,request.user,'simple_decision.commented',d.pk,{'text':text})
    elif action in ('approve','reject'):
        require_reauth(request)
        if d.status!='pending': raise BusinessError('This decision is already closed.')
        if d.votes.filter(user=request.user).exists(): raise BusinessError('You already voted on this decision.')
        DecisionVote.objects.create(decision=d,user=request.user,choice=action)
        if action=='reject': d.status='rejected'
        elif d.votes.filter(choice='approve').count()>=d.required_approvals: d.status='approved'
        d.save(update_fields=['status'])
        services.audit(div,request.user,'simple_decision.'+action,d.pk,{'status':d.status})
    else: raise BusinessError('Unknown decision action.')
    return redirect('simple_decision_detail',division_id=div.pk,decision_id=d.pk)

@workspace_view
def documents(request,div):
    if request.method=='POST':
        require_role(request.user,div,('operator',))
        if not request.FILES.getlist('files'): raise BusinessError('Choose at least one file.')
        for item in request.FILES.getlist('files'):
            doc=services.upload(div,request.user,item.name,item.content_type,item.read(10*1024*1024+1))
            DocumentLink.objects.create(document=doc,category=request.POST.get('category','')[:80],
                supplier=request.POST.get('supplier','')[:200],reference=request.POST.get('reference','')[:100])
        return redirect('simple_documents',division_id=div.pk)
    qs=Document.objects.filter(division=div).defer('content').order_by('-created_at')
    q=request.GET.get('q','').strip()[:100]
    if q: qs=qs.filter(Q(name__icontains=q)|Q(simple_links__category__icontains=q)|Q(simple_links__supplier__icontains=q)|Q(simple_links__reference__icontains=q)|Q(workflowattachment__category__icontains=q)|Q(workflowattachment__supplier__icontains=q)|Q(workflowattachment__reference__icontains=q)).distinct()
    start,end,_=filters(request)
    if start: qs=qs.filter(created_at__date__gte=start)
    if end: qs=qs.filter(created_at__date__lte=end)
    return render(request,'core/simple_documents.html',ctx(request,div,'documents',documents=qs[:200]))

@workspace_view
def document_preview(request,div,document_id):
    doc=get_object_or_404(Document,division=div,pk=document_id)
    if doc.scan_status!='clean' and not div.is_demo: raise BusinessError('Live preview requires a configured malware scanner.')
    if doc.media_type not in ('image/png','image/jpeg','application/pdf','text/plain'): raise BusinessError('Preview type is unsupported.')
    response=FileResponse(io.BytesIO(bytes(doc.content)),content_type=doc.media_type)
    response['Content-Disposition']='inline'
    response['Content-Security-Policy']="default-src 'none'; sandbox"
    return response

@workspace_view
def reports_view(request,div):
    data=summary(div,*filters(request)); legacy=reports.report(div)
    from .manufacturing import manufacturing_totals,performance_by_period
    mfg=manufacturing_totals(div,*filters(request))
    return render(request,'core/simple_reports.html',ctx(request,div,'reports',summary=data,legacy=legacy,
        mfg=mfg,performance=performance_by_period(div,*filters(request)[:2]),entries=filtered_entries(div,request)[:100]))

@workspace_view
def export_view(request,div,kind):
    entries=list(filtered_entries(div,request)); data=summary(div,*filters(request))
    from .manufacturing import manufacturing_totals,performance_by_period
    mfg=manufacturing_totals(div,*filters(request))
    performance=performance_by_period(div,*filters(request)[:2])
    def safe_cell(value):
        value=str(value)
        return "'"+value if value.lstrip().startswith(('=','+','-','@')) else value
    if kind=='csv':
        out=io.StringIO(); w=csv.writer(out); w.writerow(['Date','Money direction','Cash amount SAR','Total owed SAR','Category','Description','Payer or payee','Method','Reference','Profit treatment','Recognized amount SAR'])
        for t in entries: w.writerow([t.date,t.direction,t.amount,t.total_due or '',safe_cell(t.category),safe_cell(t.description),safe_cell(t.party),t.method,safe_cell(t.reference),t.treatment,t.recognized_amount])
        response=HttpResponse(out.getvalue(),content_type='text/csv'); response['Content-Disposition']='attachment; filename="transactions.csv"'; return response
    if kind=='xlsx':
        from openpyxl import Workbook
        wb=Workbook(); ws=wb.active; ws.title='Transactions'
        ws.append(['Date','Direction','Cash SAR','Total owed SAR','Category','Description','Payer or payee','Method','Reference','Profit treatment','Recognized SAR'])
        for t in entries: ws.append([t.date,t.direction,float(t.amount),float(t.total_due) if t.total_due is not None else None,safe_cell(t.category),safe_cell(t.description),safe_cell(t.party),t.method,safe_cell(t.reference),t.treatment,float(t.recognized_amount)])
        performance_sheet=wb.create_sheet('Manufacturing summary')
        performance_sheet.append(['Measure','SAR or units'])
        for label,key in [('Sales revenue','sales'),('Goods cost on dispatch','cogs'),('Gross profit','gross'),('Expenses and write-offs','expenses'),('Estimated profit','profit'),('Inventory value','inventory_value'),('Units in stock','inventory_units'),('Mold balances due','mold_due'),('Production balances due','production_due'),('Customer balances due','receivables'),('Capital contributed','capital'),('Capital repaid','capital_repaid'),('Confirmed funding due','funding_due')]:
            performance_sheet.append([label,float(mfg[key])])
        for period,rows in [('Monthly performance',performance['months']),('Yearly performance',performance['years'])]:
            sheet=wb.create_sheet(period); sheet.append(['Period','Sales SAR','Goods cost SAR','Expenses SAR','Other adjustments SAR','Estimated profit SAR'])
            for name,row in rows:
                sheet.append([name,*[float(row[key]) for key in ('sales','cogs','expenses','other','profit')]])
        buffer=io.BytesIO(); wb.save(buffer); buffer.seek(0)
        return FileResponse(buffer,as_attachment=True,filename='transactions.xlsx',content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    if kind=='pdf':
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        buffer=io.BytesIO(); pdf=canvas.Canvas(buffer,pagesize=A4); y=800
        pdf.setFont('Helvetica-Bold',15); pdf.drawString(45,y,'Business finance report (SAR)'); y-=28
        pdf.setFont('Helvetica',10)
        for label,key in [('Sales revenue','sales'),('Goods cost on dispatch','cogs'),('Gross profit','gross'),('Estimated accounting profit','profit'),('Inventory value','inventory_value'),('Mold payments due','mold_due'),('Production payments due','production_due'),('Customer balances due','receivables')]:
            pdf.drawString(45,y,f'{label}: {mfg[key]}'); y-=16
        y-=12
        for label,key in [('Money received','received'),('Money paid','paid'),('Net cash movement','net_cash'),('Income','income'),('Expenses','expenses'),('Estimated profit','profit'),('Receivables','receivable'),('Payables','payable')]:
            pdf.drawString(45,y,f'{label}: {data[key]}'); y-=16
        y-=16; pdf.drawString(45,y,'Recent transactions'); y-=18
        for t in entries[:35]:
            if y<45: pdf.showPage(); pdf.setFont('Helvetica',10); y=800
            line=f'{t.date}  {t.direction}  {t.amount}  {t.category[:20]}  {t.description[:45]}'
            pdf.drawString(45,y,line.encode('latin-1','replace').decode('latin-1')); y-=15
        pdf.save(); buffer.seek(0)
        return FileResponse(buffer,as_attachment=True,filename='finance-report.pdf',content_type='application/pdf')
    raise BusinessError('Unknown export format.')
